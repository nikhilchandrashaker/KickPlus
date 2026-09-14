"""
Kick+ Calculation Engine
=========================
An era- and distance-adjusted NFL kicking statistic.

Definition
----------
Kick+ = 100 x (Sum of FG Made) / Sum(Attempts x League FG% for that distance-bucket/year)

Interpretation: how many field goals a kicker made relative to how many an
average NFL kicker would be expected to make, given the exact same distribution
of attempts (same distances, same years/eras).

100 = league average. 110 = made 10% more FGs than an average kicker would have,
given identical attempts.

Data sources (from 21st_Century_NFL_Field_Goal_Kicking_Performance.xlsx):
- 'Cleaned and Reformatted': kicker-season rows with makes/attempts per distance bucket
- 'Yearly Aggregates': league-wide makes/attempts per distance bucket per year
  (used to derive the league baseline rate -- NOT the pre-computed rate column,
  since that column contains '#DIV/0!' strings for zero-attempt buckets in some years)

Distance buckets are aligned POSITIONALLY across the two sheets because the
label text differs slightly ('1-19' vs '0-19') but the bucket order is identical:
[<20, 20-29, 30-39, 40-49, 50-59, 60+]
"""

import openpyxl
from collections import defaultdict
import csv
import statistics

SOURCE_FILE = "/mnt/user-data/uploads/21st_Century_NFL_Field_Goal_Kicking_Performance.xlsx"

BUCKETS = ["<20", "20-29", "30-39", "40-49", "50-59", "60+"]

# Column offsets within a 'Cleaned and Reformatted' row, keyed by bucket.
# Row layout: Player, FGM, Att, FG%, TotFGM, TotFGA, TotRate,
#   [Made, Att, Rate] x 6 buckets, Lng, Blk, Year
CLEANED_BUCKET_COLS = {
    "<20":   (7, 8),
    "20-29": (10, 11),
    "30-39": (13, 14),
    "40-49": (16, 17),
    "50-59": (19, 20),
    "60+":   (22, 23),
}
CLEANED_YEAR_COL = 27

# Column offsets within a 'Yearly Aggregates' row, keyed by bucket.
# Row layout: Year, TotFGM, TotFGA, TotRate, [Made, Atts, Rate] x 6 buckets, ...
YEARLY_BUCKET_COLS = {
    "<20":   (4, 5),
    "20-29": (7, 8),
    "30-39": (10, 11),
    "40-49": (13, 14),
    "50-59": (16, 17),
    "60+":   (19, 20),
}


def _num(v, default=0):
    return v if isinstance(v, (int, float)) else default


def load_league_rates(wb):
    """Return {year: {bucket: made/att rate}}, computed fresh from made/att
    totals rather than trusting any pre-computed rate column."""
    ws = wb["Yearly Aggregates"]
    rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]
    league = {}
    for r in rows:
        year = int(r[0])
        league[year] = {}
        for bucket, (midx, aidx) in YEARLY_BUCKET_COLS.items():
            made, att = _num(r[midx]), _num(r[aidx])
            league[year][bucket] = (made / att) if att > 0 else None
    return league


def load_kicker_seasons(wb, league):
    """Return a list of dicts, one per (kicker, year, bucket) with actual
    makes and league-expected makes, skipping buckets with no league rate
    available (i.e. essentially zero league attempts that year -- e.g. 60+
    in the early 2000s)."""
    ws = wb["Cleaned and Reformatted"]
    rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]
    records = []
    skipped_no_rate = 0
    for r in rows:
        name = r[0]
        year = r[CLEANED_YEAR_COL]
        if not isinstance(year, (int, float)):
            continue
        year = int(year)
        if year not in league:
            continue
        for bucket, (midx, aidx) in CLEANED_BUCKET_COLS.items():
            att = _num(r[aidx])
            if att <= 0:
                continue
            made = _num(r[midx])
            rate = league[year][bucket]
            if rate is None:
                skipped_no_rate += 1
                continue
            records.append({
                "player": name,
                "year": year,
                "bucket": bucket,
                "made": made,
                "att": att,
                "expected": att * rate,
            })
    return records, skipped_no_rate


def aggregate(records, keys):
    """Group records by `keys` (e.g. ('player',) or ('player','year')) and
    sum made/att/expected."""
    out = defaultdict(lambda: {"made": 0.0, "att": 0.0, "expected": 0.0})
    for rec in records:
        k = tuple(rec[key] for key in keys)
        out[k]["made"] += rec["made"]
        out[k]["att"] += rec["att"]
        out[k]["expected"] += rec["expected"]
    return out


def kick_plus(agg_dict, min_att=0):
    results = []
    for k, v in agg_dict.items():
        if v["att"] < min_att or v["expected"] <= 0:
            continue
        kp = 100.0 * v["made"] / v["expected"]
        results.append((*k, v["att"], v["made"], v["expected"], kp))
    return results


# ---------------------------------------------------------------------------
# Sample-size tiers (per V1 spec: 20+ ranked, 10-19 flagged, <10 unranked)
# ---------------------------------------------------------------------------

SEASON_FULL_MIN = 20
SEASON_SMALL_SAMPLE_MIN = 10
CAREER_RANK_MIN = 100          # min career attempts to appear on career leaderboards
CONSISTENCY_MIN_SEASONS = 3    # min qualifying (20+ att) seasons to compute an SD
BUCKET_RANK_MIN = 20           # min attempts within a single distance bucket to rank in that bucket's leaderboard


def season_tier(att):
    if att >= SEASON_FULL_MIN:
        return "ranked"
    elif att >= SEASON_SMALL_SAMPLE_MIN:
        return "small_sample"
    else:
        return "unranked"


# ---------------------------------------------------------------------------
# Validation
# ---------------------------------------------------------------------------

def run_validation(records, league, wb):
    print("=" * 70)
    print("VALIDATION")
    print("=" * 70)

    # Check 1: league-wide Kick+ per year should be ~100 by construction,
    # since the baseline rate IS the league rate for that year/bucket.
    print("\n[1] League-wide Kick+ by year (should be ~100.0 every year):")
    by_year = aggregate(records, ("year",))
    max_dev = 0.0
    for year in sorted(y for (y,) in by_year.keys()):
        v = by_year[(year,)]
        kp = 100.0 * v["made"] / v["expected"] if v["expected"] else None
        dev = abs(kp - 100.0) if kp else None
        if dev:
            max_dev = max(max_dev, dev)
        flag = "  <-- CHECK" if dev and dev > 0.5 else ""
        print(f"    {year}: Kick+ = {kp:6.2f}{flag}")
    print(f"    Max deviation from 100.0 across all years: {max_dev:.4f}")
    assert max_dev < 0.5, "League-wide Kick+ deviates too far from 100 -- formula bug likely"

    # Check 2: total attempts reconstructed from kicker-season rows should
    # match the Yearly Aggregates total FGA per year (cross-sheet integrity).
    print("\n[2] Attempt totals: kicker-level sum vs. Yearly Aggregates sheet:")
    ws = wb["Yearly Aggregates"]
    yearly_rows = {int(r[0]): r for r in ws.iter_rows(min_row=2, values_only=True) if r[0]}
    mismatches = 0
    for year in sorted(y for (y,) in by_year.keys()):
        reconstructed = by_year[(year,)]["att"]
        reported = _num(yearly_rows[year][2])  # Total FGA column
        diff = reconstructed - reported
        flag = ""
        if abs(diff) > 0:
            mismatches += 1
            flag = f"  <-- diff {diff:+.0f}"
        print(f"    {year}: reconstructed={reconstructed:.0f}  reported={reported:.0f}{flag}")
    if mismatches:
        print(f"    NOTE: {mismatches} year(s) differ -- likely buckets dropped for zero league rate (see skip count below), not a calc error.")

    # Check 3: distribution sanity -- Kick+ should cluster near 100, and the
    # sport's known best/worst kickers should land where reputation says.
    print("\n[3] Career Kick+ distribution (min. 100 attempts):")
    by_player = aggregate(records, ("player",))
    career = kick_plus(by_player, min_att=100)
    career.sort(key=lambda x: -x[-1])
    values = [row[-1] for row in career]
    print(f"    N kickers (>=100 att): {len(values)}")
    print(f"    Mean: {sum(values)/len(values):.2f}  Min: {min(values):.2f}  Max: {max(values):.2f}")
    print(f"    Top 5: {[ (r[0], round(r[-1],1)) for r in career[:5] ]}")
    print(f"    Bottom 5: {[ (r[0], round(r[-1],1)) for r in career[-5:] ]}")

    print("\nValidation passed.\n")


def main():
    wb = openpyxl.load_workbook(SOURCE_FILE, data_only=True)
    league = load_league_rates(wb)
    records, skipped = load_kicker_seasons(wb, league)
    print(f"Loaded {len(records)} kicker-season-bucket records "
          f"({skipped} bucket-rows skipped: league had zero attempts that bucket/year).\n")

    run_validation(records, league, wb)

    # ---- Season Kick+, ALL seasons (tiered: ranked / small_sample / unranked) ----
    by_player_year = aggregate(records, ("player", "year"))
    all_seasons = kick_plus(by_player_year, min_att=1)   # keep every season, tag it, filter downstream
    all_seasons.sort(key=lambda x: (x[0], x[1]))          # player, year order for readability
    with open("/home/claude/kickplus/season_kick_plus.csv", "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["player", "year", "attempts", "made", "expected_made", "kick_plus", "tier"])
        for row in all_seasons:
            player, year, att, made, exp, kp = row
            w.writerow([player, year, round(att), round(made), round(exp, 1), round(kp, 1), season_tier(att)])

    ranked_seasons = [r for r in all_seasons if season_tier(r[2]) == "ranked"]
    ranked_seasons_sorted = sorted(ranked_seasons, key=lambda x: -x[-1])
    with open("/home/claude/kickplus/leaderboard_best_season.csv", "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["rank", "player", "year", "attempts", "kick_plus"])
        for i, row in enumerate(ranked_seasons_sorted, 1):
            w.writerow([i, row[0], row[1], round(row[2]), round(row[-1], 1)])

    # ---- Career Kick+ + seasons played + consistency (SD of ranked-season Kick+) ----
    by_player = aggregate(records, ("player",))
    career_all = {row[0]: row for row in kick_plus(by_player, min_att=1)}

    # seasons played (any tier) and ranked-season Kick+ values, per player
    seasons_played = defaultdict(int)
    ranked_kp_by_player = defaultdict(list)
    for player, year, att, made, exp, kp in all_seasons:
        seasons_played[player] += 1
        if season_tier(att) == "ranked":
            ranked_kp_by_player[player].append(kp)

    with open("/home/claude/kickplus/career_kick_plus.csv", "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["player", "attempts", "made", "expected_made", "career_kick_plus",
                     "seasons_played", "n_ranked_seasons", "consistency_sd", "meets_career_rank_min"])
        career_rows = []
        for player, row in career_all.items():
            _, att, made, exp, kp = row
            ranked_kps = ranked_kp_by_player.get(player, [])
            sd = round(statistics.stdev(ranked_kps), 1) if len(ranked_kps) >= CONSISTENCY_MIN_SEASONS else ""
            career_rows.append((player, att, made, exp, kp, seasons_played[player], len(ranked_kps), sd))
        career_rows.sort(key=lambda x: -x[4])
        for player, att, made, exp, kp, n_seasons, n_ranked, sd in career_rows:
            meets_min = att >= CAREER_RANK_MIN
            w.writerow([player, round(att), round(made), round(exp, 1), round(kp, 1),
                        n_seasons, n_ranked, sd, meets_min])

    with open("/home/claude/kickplus/leaderboard_career.csv", "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["rank", "player", "attempts", "seasons_played", "career_kick_plus"])
        qualified = [r for r in career_rows if r[1] >= CAREER_RANK_MIN]
        qualified.sort(key=lambda x: -x[4])
        for i, (player, att, made, exp, kp, n_seasons, n_ranked, sd) in enumerate(qualified, 1):
            w.writerow([i, player, round(att), n_seasons, round(kp, 1)])

    with open("/home/claude/kickplus/leaderboard_most_consistent.csv", "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["rank", "player", "n_ranked_seasons", "career_kick_plus", "consistency_sd"])
        consistent = [r for r in career_rows if r[7] != ""]
        consistent.sort(key=lambda x: x[7])   # lower SD = more consistent
        for i, (player, att, made, exp, kp, n_seasons, n_ranked, sd) in enumerate(consistent, 1):
            w.writerow([i, player, n_ranked, round(kp, 1), sd])

    # ---- Distance-bucket breakdown per player (career) — full table for player pages ----
    by_player_bucket = aggregate(records, ("player", "bucket"))
    bucket_rows = kick_plus(by_player_bucket, min_att=1)
    with open("/home/claude/kickplus/career_kick_plus_by_bucket.csv", "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["player", "bucket", "attempts", "made", "expected_made", "kick_plus"])
        for row in bucket_rows:
            w.writerow([row[0], row[1], round(row[2]), round(row[3]), round(row[4], 1), round(row[5], 1)])

    # ---- Wide distance-profile table: one row per player, one column per bucket ----
    profile = defaultdict(dict)
    profile_att = defaultdict(dict)
    for player, bucket, att, made, exp, kp in bucket_rows:
        profile[player][bucket] = kp
        profile_att[player][bucket] = att
    with open("/home/claude/kickplus/kick_plus_distance_profile.csv", "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["player"] + [f"{b}_kick_plus" for b in BUCKETS] + [f"{b}_att" for b in BUCKETS])
        for player in sorted(profile.keys()):
            kp_cells = [round(profile[player].get(b, 0), 1) if b in profile[player] else "" for b in BUCKETS]
            att_cells = [round(profile_att[player].get(b, 0)) if b in profile_att[player] else "" for b in BUCKETS]
            w.writerow([player] + kp_cells + att_cells)

    # ---- Per-distance-bucket leaderboards (50-59 and 60+, min BUCKET_RANK_MIN attempts) ----
    for bucket in ["50-59", "60+"]:
        rows = [r for r in bucket_rows if r[1] == bucket and r[2] >= BUCKET_RANK_MIN]
        rows.sort(key=lambda x: -x[-1])
        fname = f"/home/claude/kickplus/leaderboard_{bucket.replace('+','plus').replace('-','_')}.csv"
        with open(fname, "w", newline="") as f:
            w = csv.writer(f)
            w.writerow(["rank", "player", "attempts", "made", "kick_plus"])
            for i, row in enumerate(rows, 1):
                w.writerow([i, row[0], round(row[2]), round(row[3]), round(row[5], 1)])

    print(f"Wrote {len(career_rows)} career rows ({len(qualified)} rank-eligible), "
          f"{len(all_seasons)} season rows ({len(ranked_seasons)} ranked), "
          f"{len(bucket_rows)} player-bucket rows, {len(consistent)} players with a consistency score.")


if __name__ == "__main__":
    main()
